import io
import json
import docx
import pytest
from fastapi.testclient import TestClient
from unittest.mock import patch

from main import app
from db.database import init_db, list_kbs, delete_kb

client = TestClient(app)

@pytest.fixture(autouse=True)
def setup_database():
    init_db()
    yield

def test_full_pipeline_end_to_end():
    # -------------------------------------------------------------
    # STEP 1: Add SAP Knowledge Base to SQLite
    # -------------------------------------------------------------
    kb_md = (
        "# SAP SD Pricing and Conditions\n"
        "Base price condition PR00 is mandatory for all standard sales orders (OR).\n"
        "Value-added tax condition MWST calculates 19% domestic tax."
    ).encode("utf-8")

    kb_upload_res = client.post(
        "/kb/upload",
        files={"file": ("sd_pricing_rules.md", io.BytesIO(kb_md), "text/markdown")},
        data={"name": "SAP SD Pricing Rules"},
    )
    assert kb_upload_res.status_code == 200, kb_upload_res.text
    kb_id = kb_upload_res.json()["kb_id"]
    assert kb_id > 0

    # Verify in KB list
    kb_list_res = client.get("/kb/list")
    assert kb_list_res.status_code == 200
    assert any(k["id"] == kb_id for k in kb_list_res.json())

    # -------------------------------------------------------------
    # STEP 2: Developer Uploads Requirement Document (DOCX)
    # -------------------------------------------------------------
    doc = docx.Document()
    doc.add_heading("Functional Requirement: Customer Sales Order Processing", level=1)
    doc.add_paragraph("Customer PO number is strictly mandatory for sold-to party 1000.")
    doc.add_paragraph("Automatic credit check triggers whenever order total exceeds 50000 EUR.")
    doc_io = io.BytesIO()
    doc.save(doc_io)
    docx_bytes = doc_io.getvalue()

    upload_res = client.post(
        "/upload",
        files={"file": ("sales_order_requirement.docx", io.BytesIO(docx_bytes), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
    )
    assert upload_res.status_code == 200, upload_res.text
    upload_data = upload_res.json()
    req_id = upload_data["doc_id"]
    assert "Automatic credit check triggers" in upload_data["text"]


    # -------------------------------------------------------------
    # STEP 3: Gemini 2.5 Flash Decision Model Filters Relevant KBs
    # -------------------------------------------------------------
    mock_relevance_result = {
        "status": "success",
        "relevant_kbs": [{"id": kb_id, "name": "SAP SD Pricing Rules", "content": "PR00"}],
        "decision_model_used": "gemini-2.5-flash",
        "total_kbs_evaluated": 1,
        "message": "Identified 1 relevant knowledge base(s).",
    }

    with patch("routes.generate.get_relevant_kbs", return_value=mock_relevance_result):
        eval_res = client.post(
            "/generate/evaluate-kbs",
            json={"requirement_id": req_id}
        )
        assert eval_res.status_code == 200, eval_res.text
        eval_data = eval_res.json()
        assert eval_data["status"] == "success"
        assert len(eval_data["relevant_kbs"]) == 1
        assert eval_data["relevant_kbs"][0]["id"] == kb_id
        assert eval_data["decision_model_used"] == "gemini-2.5-flash"

    # -------------------------------------------------------------
    # STEP 4: Test Case Generation Agent (Gemini 1.5 Flash)
    # -------------------------------------------------------------
    mock_test_cases = [
        {
            "id": "TC-001",
            "title": "Create standard sales order with mandatory PO and valid PR00 price",
            "related_component": "SAP SD Order & Pricing",
            "type": "positive",
            "preconditions": "Sold-to party 1000 exists and PR00 condition active",
            "steps": [
                "Open transaction VA01",
                "Enter sold-to party 1000 and valid PO number PO-9988",
                "Add material line item and verify PR00 calculation",
                "Save sales order"
            ],
            "expected_result": "Sales order created successfully with document number generated.",
            "status": "pending"
        },
        {
            "id": "TC-002",
            "title": "Create sales order without mandatory Customer PO number",
            "related_component": "SAP SD Order Validation",
            "type": "negative",
            "preconditions": "Sold-to party 1000 configuration active",
            "steps": [
                "Open transaction VA01",
                "Leave Purchase Order field blank",
                "Click Save button"
            ],
            "expected_result": "System raises error: Customer PO number is strictly mandatory.",
            "status": "pending"
        },
        {
            "id": "TC-003",
            "title": "Credit check threshold boundary validation at 50,000 EUR",
            "related_component": "SAP SD Credit Check",
            "type": "boundary",
            "preconditions": "Credit limit rule set at 50000 EUR",
            "steps": [
                "Create sales order with total value exactly 50000 EUR -> check no block",
                "Create sales order with total value 50001 EUR -> check credit block applied"
            ],
            "expected_result": "Order at 50000 EUR is unblocked; order at 50001 EUR is placed on credit block.",
            "status": "pending"
        }
    ]

    with patch("services.generator.generate_test_cases", return_value=mock_test_cases):
        gen_res = client.post(
            "/generate",
            json={
                "requirement_id": req_id,
                "relevant_kb_ids": [kb_id]
            }
        )
        assert gen_res.status_code == 200, gen_res.text
        gen_data = gen_res.json()
        assert gen_data["status"] == "success"
        assert len(gen_data["test_cases"]) == 3
        types = {tc["type"] for tc in gen_data["test_cases"]}
        assert "positive" in types
        assert "negative" in types
        assert "boundary" in types

    # -------------------------------------------------------------
    # STEP 5 & 6: Human Review, Edit & Approval -> JSON Download
    # -------------------------------------------------------------
    # Human edits TC-001 title inline during review
    edited_cases = list(gen_data["test_cases"])
    edited_cases[0]["title"] = "Create standard sales order (Reviewed and Verified by Human QA)"

    approve_res = client.post(
        "/approve",
        json={
            "project_name": "SAP E2E Verification",
            "requirement_id": req_id,
            "test_cases": edited_cases,
        }
    )
    assert approve_res.status_code == 200, approve_res.text
    assert "attachment; filename=\"sap_test_cases_" in approve_res.headers.get("content-disposition", "")
    assert "application/json" in approve_res.headers.get("content-type", "")

    approved_json = json.loads(approve_res.text)
    assert approved_json["total_test_cases"] == 3
    assert approved_json["coverage_summary"] == {"positive": 1, "negative": 1, "boundary": 1}
    for tc in approved_json["test_cases"]:
        assert tc["status"] == "approved"
    assert "Reviewed and Verified by Human QA" in approved_json["test_cases"][0]["title"]

    # -------------------------------------------------------------
    # STEP 7: Single Port Web Serving Verification
    # -------------------------------------------------------------
    web_res = client.get("/")
    assert web_res.status_code == 200
    assert "SAP Test Case Generator" in web_res.text

    # Cleanup KB
    delete_kb(kb_id)

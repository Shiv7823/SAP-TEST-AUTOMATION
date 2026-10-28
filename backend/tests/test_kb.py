import io
import pytest
from fastapi.testclient import TestClient

from main import app
from db.database import init_db, insert_kb, list_kbs, get_kb_by_id, delete_kb, DB_PATH

@pytest.fixture(autouse=True)
def setup_test_db():
    """Ensure database schema is initialized before every test."""
    init_db()
    yield

def test_db_crud():
    # 1. Insert
    kb_id = insert_kb(
        name="SAP Pricing Rules KB",
        content="Pricing calculation requires condition PR00 and net value calculation.",
        file_type="markdown",
    )
    assert kb_id is not None
    assert kb_id > 0

    # 2. Get by ID
    kb = get_kb_by_id(kb_id)
    assert kb is not None
    assert kb["name"] == "SAP Pricing Rules KB"
    assert "PR00" in kb["content"]
    assert kb["file_type"] == "markdown"

    # 3. List
    kbs = list_kbs()
    assert any(item["id"] == kb_id for item in kbs)

    # 4. Delete
    deleted = delete_kb(kb_id)
    assert deleted is True

    # Confirm deleted
    assert get_kb_by_id(kb_id) is None

def test_api_kb_lifecycle():
    client = TestClient(app)

    # 1. Upload markdown KB via /kb/upload
    md_content = b"# SAP SD Shipping Rules\nGoods issue requires complete picking and packing."
    files = {
        "file": ("sd_shipping_rules.md", io.BytesIO(md_content), "text/markdown")
    }
    data = {"name": "SD Shipping Rules"}

    upload_res = client.post("/kb/upload", files=files, data=data)
    assert upload_res.status_code == 200, upload_res.text
    res_data = upload_res.json()
    assert "kb_id" in res_data
    kb_id = res_data["kb_id"]
    assert res_data["name"] == "SD Shipping Rules"

    # 2. List KBs via /kb/list
    list_res = client.get("/kb/list")
    assert list_res.status_code == 200
    kbs = list_res.json()
    matched = [item for item in kbs if item["id"] == kb_id]
    assert len(matched) == 1
    assert matched[0]["name"] == "SD Shipping Rules"
    assert matched[0]["file_type"] == "markdown"

    # 3. Delete KB via /kb/{kb_id}
    del_res = client.delete(f"/kb/{kb_id}")
    assert del_res.status_code == 200
    assert del_res.json()["kb_id"] == kb_id

    # 4. Verify no longer in list
    list_res2 = client.get("/kb/list")
    assert not any(item["id"] == kb_id for item in list_res2.json())

def test_parser_docx_and_text():
    import docx
    from services.parser import extract_text

    # Test docx parser
    doc = docx.Document()
    doc.add_heading("SAP SD Architecture", level=1)
    doc.add_paragraph("Sales Orders require order type OR and sales org 1000.")
    doc_io = io.BytesIO()
    doc.save(doc_io)
    docx_bytes = doc_io.getvalue()

    extracted_docx = extract_text(docx_bytes, "sd_arch.docx")
    assert "SAP SD Architecture" in extracted_docx
    assert "Sales Orders require order type OR" in extracted_docx

    # Test plain text / markdown parser
    txt_bytes = "Plain text requirement for SAP MM module.".encode("utf-8")
    extracted_txt = extract_text(txt_bytes, "req.txt")
    assert "Plain text requirement for SAP MM module." in extracted_txt

